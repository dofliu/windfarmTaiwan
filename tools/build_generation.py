#!/usr/bin/env python3
"""風場的實際年發電量 · Actual yearly generation of wind farms → data/global/generation.json

  python3 tools/build_generation.py uswtdb_V9_1_20260928.csv f923_2023.zip f923_2024.zip f923_2025.zip august_generator2026.xlsx \
      data/global/sources/taipower_renewable_generation_17140.csv data/global/sources/taipower_wind_stations_17141.csv

美國：EIA-923（美國能源資訊署，各電廠逐月淨發電量，公有領域）。USWTDB 每部風機都附 EIA 電廠代碼（tools/build_turbines.py 寫進
turbines.json 的 eia 欄），所以本站風場→USWTDB 機組→EIA 電廠可以直接接上。寧可少配：
  1. 一座 EIA 電廠的機組必須全部在同一座本站風場裡（電廠跨好幾座風場時，發電量分不開，不用）。
  2. 某一年只有在這些電廠的機組全部在前一年以前商轉、沒有在當年改裝（USWTDB 的 p_year、t_retro_yr）時才採用，
     避免把部分年度當成全年。
  3. 容量因數＝淨發電量 ÷（EIA-860M 登記的風力機組裝置容量 × 8,760 小時）：分子分母都是 EIA 同一座電廠的數字。
     USWTDB 有時少收一部分機組（例：Salt Fork 收 64 部、128 MW，EIA 登記 174 MW），用 USWTDB 加總當分母會把容量因數算高；
     USWTDB 與 EIA 的容量相差 10% 以上表示風場與電廠對不乾淨，整座不用。當年有機組商轉或除役（EIA-860M 的商轉年、除役年）的年份不用。
     不在 5–65% 之間的年份也不用（多為資料錯誤或停機）。
EIA-923 年度檔的網址：https://www.eia.gov/electricity/data/eia923/（最新年在 xls/，較舊的在 archive/xls/）；
EIA-860M（各機組的裝置容量、商轉與除役年月，每月更新）：https://www.eia.gov/electricity/data/eia860m/（例：xls/august_generator2026.xlsx）。
台灣：台電「自建之各類再生能源發電量」（政府資料開放平臺 17140，各發電站逐月淨發電量，2024 年起）與
「風力發電站資料」（17141，各站裝置容量），存在 data/global/sources/（下載：service.taipower.com.tw/data/opendata/apply/file/d693001/001.csv、d693002/001.csv），政府資料開放授權條款第 1 版。只有台電自有的發電站，民營風場沒有逐場的官方數字。
發電站與本站風場的對應寫在 TW_STATIONS，台電公布的裝置容量要在本站容量 ±15% 內才用（對不上多半是本站的紀錄只含其中幾期）；只用 12 個月都有數字的年份。
機型（給「發電表現」的同機型比較用，m 欄）：寧可不寫——美國取該風場 EIA 電廠在 USWTDB 的所有機組，只有一種製造商＋型號（且單機容量相同）時才寫，
並附輪轂高度與葉輪直徑；台灣取本站風場紀錄的 turbine 欄，只有一種機型（沒有「+」、寫得出製造商）、而且台電發電站容量與本站紀錄相差 3% 以內
（相差較多時發電站可能還含其他機組）才寫。
風場資料或 USWTDB 重建後要重跑（先跑 tools/build_turbines.py）。

US: EIA-923 (U.S. Energy Information Administration, monthly net generation by plant, public domain). USWTDB gives every turbine its
EIA plant code (written by tools/build_turbines.py as the "eia" field of turbines.json), so site farm → USWTDB turbines → EIA plant
links directly. Conservative: an EIA plant is used only when all of its turbines are in one site farm; a year is used only when every
turbine of those plants started before that year and was not retrofitted in it; capacity factor = net generation / (EIA-860M nameplate
capacity of the plants' wind generators × 8,760 h), so numerator and denominator are the same agency's figures for the same plant
(USWTDB sometimes lacks part of a plant, e.g. Salt Fork: 64 turbines, 128 MW in USWTDB, 174 MW at EIA, which inflated the factor). Farms
whose USWTDB and EIA capacities differ by 10% or more are dropped, as are years in which a generator entered service or retired, and
years outside 5–65%. Re-run after rebuilding the farm layer or USWTDB (run
tools/build_turbines.py first).
Taiwan: Taipower's own stations only (open data 17140, monthly net generation by station from 2024; 17141, station capacity; Taiwan
Open Government Data Licence v1). Stations map to site farms in TW_STATIONS and are used only when Taipower's capacity is within
±15% of the site's record; only years with all 12 months reported. Private farms have no official per-farm figures.
Turbine model (field m, for the same-model comparison in the Output dialog), conservative: in the US, from every USWTDB turbine of the
farm's EIA plants, written only when they share one manufacturer + model and unit rating (with hub height and rotor diameter); in Taiwan,
from the site record's turbine field, only when it names a single model with its manufacturer and Taipower's station capacity is within
3% of the site record (a larger gap means the station may include other machines).
"""
import collections
import csv
import json
import re
import sys
import zipfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
TURB = ROOT / 'data/global/turbines.json'
OUT = ROOT / 'data/global/generation.json'
CF_MIN, CF_MAX = 0.05, 0.65
TW_URL = 'https://data.gov.tw/dataset/17140'
# 台電發電站（17140／17141 的中文名稱前段）→ 本站風場名稱
TW_STATIONS = {'石門風力': 'Shimen', '林口風力': 'Linkou', '蘆竹風力': 'Taoyuan Luzhu', '觀園風力': 'Dayuan Guanyuan', '大潭風力': 'Datan (Tatan)',
               '新竹香山風力': 'Hsinchu Xiangshan', '台中港區風力': 'Taichung Port', '台中電廠風力': 'Taichung Power Plant',
               '彰工風力': 'Changgong (Changbin Xianxi)', '彰化王功風力': 'Wanggong', '彰化永興風力': 'Yongxing (Fangyuan)',
               '雲林麥寮風力': 'Yunmai (Mailiao)', '雲林台西風力': 'Yunlin Taixi', '雲林四湖風力': 'Sihu', '恆春風力': 'Hengchun (Maanshan NPP)',
               '澎湖湖西風力': 'Penghu Huxi', '金沙風力': 'Kinmen Jinsha', '澎湖龍門風力': 'Penghu Longmen', '中屯風力': 'Penghu Zhongtun',
               '離岸一期風力': 'Taipower Offshore Phase 1 (Changhua)'}


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def eia923_wind(zip_path):
    """{EIA 電廠代碼: 全年淨發電量 MWh}（只取燃料代碼 WND 的列），年份"""
    with zipfile.ZipFile(zip_path) as z:
        name = next(n for n in z.namelist() if re.search(r'Schedules_2_3_4_5.*\.xlsx$', n))
        with z.open(name) as fh:
            wb = openpyxl.load_workbook(fh, read_only=True)
            ws = wb['Page 1 Generation and Fuel Data']
            head, out, year = None, collections.defaultdict(float), None
            for r in ws.iter_rows(values_only=True):
                if head is None:
                    if r and r[0] == 'Plant Id':
                        head = [str(c or '').replace('\n', ' ') for c in r]
                        ip, ifu = head.index('Plant Id'), head.index('Reported Fuel Type Code')
                        ig, iy = head.index('Net Generation (Megawatthours)'), head.index('YEAR')
                    continue
                if r[ifu] == 'WND' and r[ip]:
                    out[int(r[ip])] += num(r[ig]) or 0
                    year = year or int(r[iy])
    return dict(out), year


def eia860m_wind(xlsx):
    """{EIA 電廠代碼: [(裝置容量 MW, 商轉年, 除役年或 0)]}（只取能源代碼 WND；營運中與已除役兩張表）"""
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    out = collections.defaultdict(list)
    for sheet, retired in (('Operating', False), ('Retired', True)):
        head = None
        for r in wb[sheet].iter_rows(values_only=True):
            if head is None:
                if r and r[0] == 'Entity ID':
                    head = {k: i for i, k in enumerate(r)}
                continue
            if not r or r[head['Energy Source Code']] != 'WND' or num(r[head['Plant ID']]) is None:
                continue
            out[int(num(r[head['Plant ID']]))].append((num(r[head['Nameplate Capacity (MW)']]) or 0, int(num(r[head['Operating Year']]) or 0),
                                                        int(num(r[head['Retirement Year']]) or 0) if retired else 0))
    return out


def tw_model(turbine):
    """「Vestas V80 2.0 MW x40 (23 in 2007, 17 in 2011)」→「Vestas V80 2.0 MW」；混合機型、沒寫製造商的 → None"""
    t = (turbine or '').strip()
    if not t or '+' in t:
        return None
    t = re.sub(r'^\d+\s*[x×]\s*', '', t)                    # 「21 x Hitachi HTW5.2-127」
    t = re.sub(r'\s*[x×]\s*\d+.*$', '', t).strip()           # 「… x40 (…)」
    return t if re.match(r'[A-Za-z]', t) else None


def tw_station(name):
    """「觀園風力發電站/Guanyuan …」→「觀園風力」"""
    return re.sub(r'(發電站)?\s*/?\s*[A-Za-z].*$', '', name).strip()


def taiwan(gen_csv, cap_csv, farms):
    cap = collections.defaultdict(float)
    for r in list(csv.reader(open(cap_csv, encoding='utf-8-sig')))[1:]:
        if '風' in r[1]:
            cap[tw_station(r[2])] += num(r[6]) or 0
    mwh, months = collections.defaultdict(float), collections.defaultdict(set)
    for r in list(csv.reader(open(gen_csv, encoding='utf-8-sig')))[1:]:
        m = re.match(r'^[\d,]+(\.\d+)?', r[4].strip())            # 「-」是沒有資料；少數格子後面多了雜字
        if '風' in r[2] and m:
            k = (tw_station(r[3]), int(r[0]))
            mwh[k] += float(m.group(0).replace(',', '')) / 1000
            months[k].add(int(r[1]))
    by_name = {f['name']: f for f in farms if f['iso'] == 'TWN'}
    out, skipped = {}, []
    for st, name in TW_STATIONS.items():
        f, kw = by_name.get(name), cap.get(st)
        if f is None:
            raise SystemExit(f'TW_STATIONS: site farm {name!r} not found (renamed?)')
        if not kw or abs(kw / 1000 - f['mw']) > 0.15 * f['mw']:
            skipped.append(f'{st} {kw and kw / 1000} MW vs {name} {f["mw"]} MW')
            continue
        ys = {}
        for (s2, y), v in mwh.items():
            if s2 == st and len(months[(s2, y)]) == 12 and y > (f['year'] or 0):
                cf = v / (kw / 1000 * 8760)
                if CF_MIN <= cf <= CF_MAX:
                    ys[y] = [round(v / 1000, 1), round(cf * 100, 1)]
        if ys:
            out['TWN|' + name] = {'y': dict(sorted(ys.items())), 'mw': round(kw / 1000, 1), 'st': st}
            m = tw_model(f.get('turbine'))
            if m and abs(kw / 1000 - f['mw']) <= 0.03 * f['mw']:
                out['TWN|' + name]['m'] = m
    print(f'Taiwan: {len(out)} Taipower stations; capacity mismatch skipped: ' + '; '.join(skipped), flush=True)
    return out


def main(uswtdb_csv, *files):
    zips = [z for z in files if z.endswith('.zip')]
    e860 = eia860m_wind(next(x for x in files if x.endswith('.xlsx') and 'generator' in x))
    tw = [c for c in files if c.endswith('.csv')]
    tb = json.loads(TURB.read_text(encoding='utf-8'))
    rows = list(csv.DictReader(open(uswtdb_csv, encoding='utf-8')))
    plant_n, plant_kw, plant_start = collections.Counter(), collections.defaultdict(float), collections.defaultdict(int)
    plant_tb = collections.defaultdict(list)                    # 機型：(製造商 型號, 單機 kW, 輪轂高度, 葉輪直徑)
    for r in rows:
        e = num(r['eia_id'])
        if not e or e <= 0:
            continue
        e = int(e)
        plant_n[e] += 1
        plant_kw[e] += num(r['t_cap']) or 0
        plant_tb[e].append((' '.join(x for x in (r['t_manu'], r['t_model']) if x and x != 'missing'), num(r['t_cap']), num(r['t_hh']), num(r['t_rd'])))
        # 最後一部機組的商轉年或改裝年：這一年以前（含）都不是完整的一年
        plant_start[e] = max(plant_start[e], int(num(r['p_year']) or 0), int(num(r['t_retro_yr']) or 0))
    gen = {}
    for z in zips:
        g, y = eia923_wind(z)
        gen[y] = g
        print(f'EIA-923 {y}: {len(g):,} wind plants, {sum(g.values()) / 1e6:,.1f} TWh', flush=True)
    years = sorted(gen)

    out, skipped_shared, skipped_cf, skipped_cap = {}, 0, 0, 0
    for name, v in tb['farms'].items():
        ids = [i for i, c in v.get('eia') or []]
        if not ids:
            continue
        if any(plant_n[i] != c for i, c in v['eia']):        # 電廠有機組在別的風場（或沒對到的專案）
            skipped_shared += 1
            continue
        mw_tb = sum(plant_kw[i] for i in ids) / 1000
        units = [u for i in ids for u in e860.get(i, [])]
        start = max(plant_start[i] for i in ids)
        ys, mw = {}, None
        for y in years:
            if y <= start or not all(i in gen[y] for i in ids):
                continue
            if any(oy == y or ry == y for _, oy, ry in units):          # 當年有機組商轉或除役：不是完整一年
                continue
            cap = sum(m for m, oy, ry in units if oy < y and (not ry or ry > y))
            if not cap or abs(mw_tb - cap) >= 0.10 * cap:             # 風場的 USWTDB 機組與 EIA 電廠對不乾淨
                skipped_cap += 1
                continue
            mw = cap
            mwh = sum(gen[y][i] for i in ids)
            cf = mwh / (cap * 8760)
            if CF_MIN <= cf <= CF_MAX:
                ys[y] = [round(mwh / 1000, 1), round(cf * 100, 1)]    # [GWh, 容量因數 %]
            else:
                skipped_cf += 1
        if ys:
            out['USA|' + name] = {'y': ys, 'mw': round(mw, 1), 'eia': ids}   # mw＝最近一年的 EIA 登記容量
            tbs = [t for i in ids for t in plant_tb[i]]
            if len({(t[0], t[1]) for t in tbs}) == 1 and ' ' in tbs[0][0] and tbs[0][1]:     # 一種型號、一種單機容量
                hh = sorted(t[2] for t in tbs if t[2] and t[2] > 0)
                out['USA|' + name].update(m=tbs[0][0], hh=round(hh[len(hh) // 2]) if hh else None, rd=round(tbs[0][3]) if tbs[0][3] and tbs[0][3] > 0 else None)
    fj = json.loads((ROOT / 'data/global/wind_farms.json').read_text(encoding='utf-8'))
    farms = [dict(zip(fj['meta']['cols'], r)) for r in fj['rows']]
    if tw:
        gen_csv = next(c for c in tw if '17140' in c or 'd693001' in c)
        cap_csv = next(c for c in tw if '17141' in c or 'd693002' in c)
        out.update(taiwan(gen_csv, cap_csv, farms))
    meta = {
        'source': {'USA': 'U.S. Energy Information Administration, Form EIA-923 (net generation by plant), years ' + ', '.join(map(str, years)) +
                          '; capacity from EIA-860M (Preliminary Monthly Electric Generator Inventory)',
                   'TWN': 'Taiwan Power Company, net generation of its own renewable stations (data.gov.tw 17140) and station capacity (17141)'},
        'url': {'USA': 'https://www.eia.gov/electricity/data/eia923/', 'TWN': TW_URL},
        'license': {'USA': 'Public domain (U.S. Government work)', 'TWN': 'Open Government Data License, version 1.0 (Taiwan)'},
        'farms': len(out), 'by_iso': dict(collections.Counter(k.split('|')[0] for k in out)), 'years': years,
        'note': 'Keys are "ISO|farm name". y = {year: [net generation GWh, capacity factor %]}; mw = capacity used for the capacity factor '
                '(US: EIA-860M nameplate of the plants\' wind generators in the latest year; Taiwan: Taipower station capacity); '
                'only full years (all turbines in service before the year) of plants wholly inside one farm. m = turbine model, only when the farm has a single model '
                '(US: USWTDB manufacturer + model; hh = median hub height m, rd = rotor diameter m; Taiwan: site record, station capacity within 3%).',
    }
    OUT.write_text(json.dumps({'meta': meta, 'farms': dict(sorted(out.items()))}, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print(f'{len(out):,} farms with generation; skipped {skipped_shared} farms sharing an EIA plant, {skipped_cap} farm-years whose USWTDB and '
          f'EIA-860M capacities differ by 10%+, {skipped_cf} farm-years outside '
          f'{CF_MIN:.0%}–{CF_MAX:.0%}; {OUT.stat().st_size / 1e3:.0f} kB')


if __name__ == '__main__':
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    main(sys.argv[1], *sys.argv[2:])
