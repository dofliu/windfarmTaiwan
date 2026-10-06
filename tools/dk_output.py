#!/usr/bin/env python3
"""丹麥：每部風機與每座風場的實測發電量 · Denmark: measured output per turbine and per farm (Danish Energy Agency register)

由 tools/build_generation.py 呼叫（輸入兩個檔時）：
  Vinddata（Stamdataregister for vindkraftanlæg：每部風機的容量、葉輪、輪轂、廠牌型號、陸域／海域、UTM 座標與逐月發電量 kWh）
  Parkproduktion（整場一起計量的風場的逐月發電量，以 Parknummer 連到風機）
下載：https://ens.dk/analyser-og-statistik/data-oversigt-over-energisektoren（「Vinddata」與「Parkproduktion」，約每 2 個月更新）。
授權：丹麥能源署（Energistyrelsen）的資料使用條款——免費、非專屬、可複製散布與修改，須標示 Energistyrelsen、資料集名稱與取用時間
（https://dataforsyningen.dk/asset/PDF/rettigheder_vilkaar/Energistyrelsen%20-%20Vilk%C3%A5r%20for%20brug%20af%20data.pdf）。
注意：丹麥能源署基於個資，只公布公司持有的風機的發電量（個人、獨資與合夥 I/S 持有的沒有，混合電廠也沒有）；整場計量的風場只有合計值，分不到單部風機。

規則（寧可少列）：
  1. 某一年只有在機組全部在前一年以前併網、當年沒有除役、12 個月都有數字時才採用；容量因數 5–65% 以外的不用（多為資料錯誤）。
  2. 單部風機：只列有自己逐月發電量的（不含整場計量的風場成員）。
  3. 風場：整場計量的風場（Parknummer）與有發電量的單機，依位置歸到最近的本站丹麥營運中風場（海域 12 km、陸域 3 km 內，陸域／海域要相同，
     併網年份不早於該風場商轉年的前一年——更早的是附近別的風機），歸到同一座本站風場的機組容量加總要在本站紀錄 ±15% 內才寫進該風場
     （差太多表示有沒公布發電量的機組，或對錯風場）。
  4. 機型（同機型比較用）：分組依廠牌＋葉輪直徑＋單機容量（登記檔型號寫法不一）。風場要所有風機同一組、型號寫法也只有一種才寫，
     本站紀錄寫明機型時用本站的寫法；單部風機的顯示名稱取該組在登記檔最常見的寫法。
註：登記檔「Kommune」與「Type af placering」兩欄的標題與內容互換，這裡依內容判斷（LAND／HAV）。

Called by tools/build_generation.py with the "Vinddata" (per-turbine master data and monthly kWh) and "Parkproduktion" (monthly kWh of
farms metered as a whole, linked by park number) workbooks from the Danish Energy Agency. Licence: the agency's data terms (free reuse;
credit Energistyrelsen, the dataset name and the retrieval date). Production is only published for company-owned turbines (privacy: none for
private persons, sole proprietorships, partnerships (I/S) or hybrid parks), and
farms metered as a whole have only a total. Rules (conservative): a year counts only if every turbine was connected before it, none was
decommissioned in it and all 12 months are reported; capacity factors outside 5–65% are dropped; farms take parks and producing single
turbines within 12 km (sea) / 3 km (land) of the nearest operating site farm of the same kind (not connected more than a year before the
farm's start year, which marks an older neighbour), and are written only when the assigned
capacity is within ±15% of the site record. Models are grouped by make + rotor diameter + unit rating (the register spells one model
several ways); a farm gets a model only when all its turbines fall in one group with one spelling.
"""
import collections
import datetime as dt
import math
import re

import openpyxl

YEARS = (2023, 2024, 2025)
CF_MIN, CF_MAX = 0.05, 0.65
SRC = {'name': 'Energistyrelsen, Stamdataregister for vindkraftanlæg (Vinddata, Parkproduktion)',
       'url': 'https://ens.dk/analyser-og-statistik/data-oversigt-over-energisektoren',
       'terms': 'https://dataforsyningen.dk/asset/PDF/rettigheder_vilkaar/Energistyrelsen%20-%20Vilk%C3%A5r%20for%20brug%20af%20data.pdf'}


def num(v):
    if v in (None, ''):
        return None
    try:
        return float(str(v).strip().replace('.', '').replace(',', '.')) if isinstance(v, str) and ',' in v else float(v)
    except (TypeError, ValueError):
        return None


def utm32_to_wgs84(x, y):
    """ETRS89 / UTM zone 32N（EPSG:25832）→ 經緯度；Krüger 級數，誤差遠小於 1 m。"""
    a, f = 6378137.0, 1 / 298.257222101
    k0, lon0, x0 = 0.9996, math.radians(9), 500000.0
    n = f / (2 - f)
    A = a / (1 + n) * (1 + n * n / 4 + n ** 4 / 64)
    b1 = n / 2 - 2 * n * n / 3 + 37 * n ** 3 / 96
    b2 = n * n / 48 + n ** 3 / 15
    b3 = 17 * n ** 3 / 480
    d1 = 2 * n - 2 * n * n / 3 - 2 * n ** 3
    d2 = 7 * n * n / 3 - 8 * n ** 3 / 5
    d3 = 56 * n ** 3 / 15
    xi, eta = y / (k0 * A), (x - x0) / (k0 * A)
    xi_ = xi - sum(b * math.sin(2 * j * xi) * math.cosh(2 * j * eta) for j, b in ((1, b1), (2, b2), (3, b3)))
    eta_ = eta - sum(b * math.cos(2 * j * xi) * math.sinh(2 * j * eta) for j, b in ((1, b1), (2, b2), (3, b3)))
    chi = math.asin(math.sin(xi_) / math.cosh(eta_))
    lat = chi + sum(d * math.sin(2 * j * chi) for j, d in ((1, d1), (2, d2), (3, d3)))
    lon = lon0 + math.atan(math.sinh(eta_) / math.cos(xi_))
    return math.degrees(lat), math.degrees(lon)


def km(a, b):
    la1, lo1, la2, lo2 = map(math.radians, (a[0], a[1], b[0], b[1]))
    h = math.sin((la2 - la1) / 2) ** 2 + math.cos(la1) * math.cos(la2) * math.sin((lo2 - lo1) / 2) ** 2
    return 2 * 6371 * math.asin(math.sqrt(h))


def date(v):
    s = str(v or '').strip()[:10]
    try:
        return dt.date.fromisoformat(s)
    except ValueError:
        return None


def read_turbines(path):
    rows = list(openpyxl.load_workbook(path, read_only=True).worksheets[0].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows[:10]) if r and r[0] and 'GSRN' in str(r[0]))
    head = rows[hi]
    months = {h: i for i, h in enumerate(head) if isinstance(h, str) and len(h) == 7 and h[4] == '-'}
    out = []
    for r in rows[hi + 1:]:
        if not r or not r[0]:
            continue
        place = next((v for v in (r[9], r[10]) if v in ('LAND', 'HAV')), '')
        muni = next((str(v) for v in (r[9], r[10]) if v and v not in ('LAND', 'HAV')), '')
        x, y = num(r[13]), num(r[14])
        ll = utm32_to_wgs84(x, y) if x and y else None
        out.append({'id': str(r[0]), 'park': str(r[1] or '').strip(), 'on': date(r[2]), 'off': date(r[3]), 'kw': num(r[4]) or 0,
                    'rd': num(r[5]), 'hh': num(r[6]), 'make': str(r[7] or '').strip(), 'model': str(r[8] or '').strip(),
                    'sea': place == 'HAV', 'land': place == 'LAND', 'muni': muni.split(' - ', 1)[-1].strip(), 'll': ll,
                    'kwh': {m: num(r[i]) for m, i in months.items()}})
    return out


def read_parks(path):
    rows = list(openpyxl.load_workbook(path, read_only=True).worksheets[0].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows[:10]) if r and r[0] and 'Parknummer' in str(r[0]))
    months = {h: i for i, h in enumerate(rows[hi]) if isinstance(h, str) and len(h) == 7 and h[4] == '-'}
    return {str(r[0]).strip(): {m: num(r[i]) for m, i in months.items()} for r in rows[hi + 1:] if r and r[0]}


def year_kwh(kwh, y):
    v = [kwh.get(f'{y}-{m:02d}') for m in range(1, 13)]
    return None if any(x is None for x in v) else sum(v)


def full_year(ts, y):
    return all(t['on'] and t['on'] < dt.date(y, 1, 1) and not (t['off'] and t['off'] <= dt.date(y, 12, 31)) for t in ts)


def units(turbines, parks):
    """發電單位：整場計量的風場（成員＋合計發電量）與有自己發電量的單機。"""
    by_park = collections.defaultdict(list)
    for t in turbines:
        if t['park']:
            by_park[t['park']].append(t)
    out = []
    for pk, ts in by_park.items():
        if pk in parks:
            out.append({'key': 'park:' + pk, 'ts': ts, 'kwh': parks[pk]})
    for t in turbines:
        if not t['park'] and any(v is not None for v in t['kwh'].values()):
            out.append({'key': t['id'], 'ts': [t], 'kwh': t['kwh']})
    for u in out:
        live = [t for t in u['ts'] if not t['off']] or u['ts']
        pts = [t['ll'] for t in live if t['ll']]
        u['ll'] = (sum(p[0] for p in pts) / len(pts), sum(p[1] for p in pts) / len(pts)) if pts else None
        u['sea'] = any(t['sea'] for t in u['ts'])
        u['kw'] = sum(t['kw'] for t in live)
        ys = {}
        for y in YEARS:
            e = year_kwh(u['kwh'], y)
            if e is None or not full_year(u['ts'], y) or not u['kw']:
                continue
            cf = e / (u['kw'] * 8760)
            if CF_MIN <= cf <= CF_MAX:
                ys[y] = [round(e / 1e6, 2), round(cf * 100, 1)]          # [GWh, 容量因數 %]
        u['y'] = ys
    return out


BRANDS = {'vestas wind systems a/s': 'Vestas', 'siemens': 'Siemens', 'siemens wind power a/s': 'Siemens',
          'siemens gamesa renewable energy a/s': 'Siemens Gamesa', 'mhi vestas offshore wind a/s': 'MHI Vestas', 'bonus': 'Bonus'}


def brand(make):
    m = ' '.join(str(make or '').split())
    if not m or m.lower() == 'ukendt':
        return None
    if m.lower() in BRANDS:
        return BRANDS[m.lower()]
    for suf in (' A/S', ' Aps', ' ApS', ' Ltd', ' GmbH', ' Limited'):
        if m.endswith(suf):
            m = m[:-len(suf)]
    return m.title() if m.isupper() else m


def model_name(t):
    """登記檔的廠牌＋型號（廠牌改成簡稱）；型號「Ukendt」（未知）或空白 → None"""
    b, md = brand(t['make']), ' '.join(str(t['model'] or '').split())
    if not b or not md or md.lower() == 'ukendt':
        return None
    return b + ' ' + md


def model_key(name):
    """同一機型的不同寫法歸成一組（V 90-3 MW＝V 90 3、SWT 2.3-93＝SWT-2.3-93）：去空白、連字號與 MW"""
    return None if not name else name.lower().replace(' mw', '').replace('mw', '').replace(' ', '').replace('-', '')


def model_of(ts):
    names = {model_name(t) for t in ts}
    keys = {model_key(n) for n in names}
    if None in names or len(keys) != 1:
        return None
    return sorted(names)[0]


BRAND_KEY = {'Siemens Gamesa': 'Siemens', 'MHI Vestas': 'Vestas'}     # 同一系列機型改了公司名稱


def gkey(t):
    """同機型比較的分組：廠牌＋葉輪直徑＋單機容量。登記檔的型號寫法不一（V 47-660／V 47、V 90-3 MW／V 90 3），用規格分組比較可靠；
    廠牌或型號不詳、沒有葉輪直徑的不分組。"""
    b = brand(t['make'])
    if not b or not model_name(t) or not t['rd'] or not t['kw']:
        return None
    return f"{BRAND_KEY.get(b, b)}|{round(t['rd'])}|{round(t['kw'])}"


def key_names(ts):
    """每個分組的顯示名稱：登記檔裡該組最常見的寫法（同數時取字母序在前的）"""
    c = collections.defaultdict(collections.Counter)
    for t in ts:
        k = gkey(t)
        if k:
            c[k][model_name(t)] += 1
    return {k: sorted(v.items(), key=lambda x: (-x[1], x[0]))[0][0] for k, v in c.items()}


def site_model(f):
    """本站風場紀錄寫明單一機型時用它（與 build_generation.py 的 tw_model 同一規則）"""
    t = str(f.get('turbine') or '').strip()
    if not t or '+' in t or 'mixed' in t.lower():
        return None
    t = re.sub(r'^\d+\s*[x×]\s*', '', t)
    t = re.sub(r'\s*[x×]\s*\d+.*$', '', t).strip()
    return t if re.match(r'[A-Za-z]', t) else None


def fits(u, f):
    """機組最新的風機併網年份不早於本站紀錄商轉年的前一年（更早的是附近別的風機，不是這座風場的）"""
    y = f.get('year')
    return not y or max((t['on'].year for t in u['ts'] if t['on']), default=0) >= y - 1


def dk_generation(vind_xlsx, park_xlsx, farms):
    """回傳 (generation.json 的丹麥風場, 單部風機清單, 統計)。farms：wind_farms.json 的列（dict）。"""
    ts, parks = read_turbines(vind_xlsx), read_parks(park_xlsx)
    us = units(ts, parks)
    site = [f for f in farms if f['iso'] == 'DNK' and f['st'] == 0 and f['lat'] is not None]
    assign = collections.defaultdict(list)
    for u in us:
        if not u['ll'] or not u['y']:
            continue
        cands = [(km(u['ll'], (f['lat'], f['lon'])), f) for f in site if (f['type'] != 0) == u['sea'] and fits(u, f)]
        cands = [c for c in cands if c[0] <= (12 if u['sea'] else 3)]
        if cands:
            assign[min(cands, key=lambda c: c[0])[1]['name']].append(u)
    by_name = {f['name']: f for f in site}
    gen, skipped = {}, []
    for name, ul in assign.items():
        f = by_name[name]
        kw = sum(u['kw'] for u in ul)
        if abs(kw / 1000 - f['mw']) > 0.15 * f['mw']:
            skipped.append(f'{name}: {kw / 1000:.1f} MW assigned vs {f["mw"]} MW')
            continue
        ys = {}
        for y in YEARS:
            if all(y in u['y'] for u in ul):                       # 歸到這座風場的每個單位那一年都要有數字
                e = sum(u['y'][y][0] for u in ul)
                ys[str(y)] = [round(e, 1), round(e * 1e6 / (kw * 8760) * 100, 1)]
        if not ys:
            continue
        tall = [t for u in ul for t in u['ts'] if not t['off']]
        hh = sorted(t['hh'] for t in tall if t['hh'])
        rd = [t['rd'] for t in tall if t['rd']]
        g = {'y': ys, 'mw': round(kw / 1000, 1), 'n': len(tall)}
        m, mk = model_of(tall), {gkey(t) for t in tall}
        if m and len(mk) == 1 and None not in mk:                   # 型號寫法與規格（廠牌、葉輪、單機容量）都只有一種
            m = site_model(f) or m                                  # 本站紀錄寫明的機型優先（登記檔偶有誤植，例：Vesterhav）
            g.update(m=m, mk=mk.pop(), hh=round(hh[len(hh) // 2]) if hh else None, rd=round(max(rd)) if rd else None)
        gen['DNK|' + name] = g
    singles, names = [], key_names(ts)
    for u in us:
        if u['key'].startswith('park:') or not u['y'] or not u['ll']:
            continue
        t = u['ts'][0]
        k = gkey(t)
        singles.append({'id': t['id'], 'll': [round(u['ll'][0], 5), round(u['ll'][1], 5)], 'kw': t['kw'], 'rd': t['rd'], 'hh': t['hh'],
                        'm': names.get(k), 'mk': k, 'sea': t['sea'], 'muni': t['muni'], 'yr': t['on'].year if t['on'] else None,
                        'y': {str(y): v for y, v in u['y'].items()}})
    months = sorted({m for t in ts for m, v in t['kwh'].items() if v is not None} | {m for k in parks.values() for m, v in k.items() if v is not None})
    stats = {'turbines': len(ts), 'parks': len(parks), 'units_with_years': sum(1 for u in us if u['y']), 'farms': len(gen),
             'singles': len(singles), 'skipped': skipped, 'through': months[-1] if months else None}     # through＝逐月資料的最後一個月
    return gen, singles, stats
